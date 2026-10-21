from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

import argparse

_parser = argparse.ArgumentParser(description='Rebuild the authored Word draft; later manual edits are not imported.')
_parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parents[1] / 'tmp' / 'generated-documents')
_args = _parser.parse_args()
OUT = _args.output.resolve()
OUT.mkdir(exist_ok=True)
NAVY, TEAL, GRAY = '18354A', '16746D', '586773'

def make_doc(title, subtitle):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.62)
    sec.left_margin = sec.right_margin = Inches(.72)
    sec.header_distance = sec.footer_distance = Inches(.28)
    normal = d.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [('Title', 28, NAVY), ('Heading 1', 19, NAVY), ('Heading 2', 12.5, TEAL), ('Heading 3', 11, NAVY)]:
        s = d.styles[name]
        s.font.name = 'Calibri'
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(color)
        s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        s.paragraph_format.space_before = Pt(9)
        s.paragraph_format.space_after = Pt(5)
    d.core_properties.title = title
    d.core_properties.subject = subtitle
    d.core_properties.author = 'Guardians | preparation draft'
    d.core_properties.comments = 'AI-assisted preparation draft. Fictional additions and proposed session conventions are labeled.'
    p = sec.header.paragraphs[0]
    r = p.add_run('GUARDIANS  /  LEAD RPG  /  UNDER THE GREEN LIGHT')
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('准备草案 · 21 OCT 2026  |  ')
    r.font.size = Pt(8)
    f = OxmlElement('w:fldSimple')
    f.set(qn('w:instr'), 'PAGE')
    p._p.append(f)
    return d

def p(d, text='', bold=False):
    x = d.add_paragraph()
    r = x.add_run(text)
    r.bold = bold
    return x

def h(d, text, level=2):
    return d.add_heading(text, level)

def bullet(d, text):
    x = d.add_paragraph(text, style='List Bullet')
    x.paragraph_format.space_after = Pt(4)
    return x

def say(d, text):
    x = d.add_paragraph()
    x.paragraph_format.left_indent = Inches(.15)
    x.paragraph_format.right_indent = Inches(.10)
    x.paragraph_format.space_before = Pt(3)
    x.paragraph_format.space_after = Pt(7)
    x.add_run('READ ALOUD  ').bold = True
    x.add_run(text)
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'EFF5F4')
    x._p.get_or_add_pPr().append(shade)

def page(d, title, subtitle=''):
    if len(d.paragraphs):
        d.add_page_break()
    h(d, title, 1)
    if subtitle:
        p(d, subtitle)

def outcomes(d, lines):
    for name, text in lines:
        x = d.add_paragraph()
        x.add_run(name + '：').bold = True
        x.add_run(text)

gm = make_doc('绿灯之下 / Under the Green Light', 'Co-GM bilingual scenario and facilitation guide')
gm.add_heading('绿灯之下', 0)
p(gm, 'UNDER THE GREEN LIGHT', True)
p(gm, 'Heliora Prime · Luminis Sea · Guardians Guild')
p(gm, '中文主持说明 + 英文朗读台词｜2 位 Co-GM + 默认 6 位玩家｜完整 45 分钟 + 复盘 10 分钟｜另附 20 分钟简版')
h(gm, '一页开团速查')
p(gm, 'The Shell 外围潮汐泵站显示绿灯，现场却在进水。十二名维修人员被困。公会要在下一轮暴潮前建立通路、转移人员、隔离泵站危险，并决定怎样如实报告一个未经验证的“安全”状态。')
bullet(gm, '读之前：只把 Player Pack 发给玩家；本手册为 GM 信息。纸质材料、D20、笔、2 枚应急物资标记即可。')
bullet(gm, '课程机制：D20、GM 设置 DC、领导力 +3 / 工程技能 +2 / 工具 +1、五档叙事结果。下文其他数值均为本草案的场次约定。')
bullet(gm, '本场约定：每骰只选一种最相关加成；准备优势与创意降难度取较高者，合计最多降 DC 2；自然 1 不另加灾难。教师已有明确规定时以教师规定为准。')
bullet(gm, '只记录：2 份应急物资；通路、转移、隔离各一个准备优势勾；三个执行结果。无 HP、无 Boss、无个人淘汰。')
h(gm, '六段节奏')
for t in ['0–4 分：开场、目标、规则；发公共任务单。', '4–13 分：发三张证据卡，90 秒拼出故障解释；三次准备检定。', '13–20 分：三名 NPC 协商；确认人员、医疗和报告责任。', '20–24 分：选择崖壁、海上、泄洪或自创混合路线；公开代价。', '24–40 分：三次团队执行检定，分别解决通路、转移、隔离。', '40–45 分：逐项叙述结局，提交口头事故简报；45–55 分复盘，最后留时间评价。']:
    bullet(gm, t)
h(gm, '两位 GM 的分工')
p(gm, 'GM A：开场、Mira 和 Tavi 的台词、组织讨论。GM B：Dr. Vale、出示卡片、记录修正值及后果、计时。执行段互换主导。游戏前写下六名玩家：准备①____②____③____；执行①____②____③____。六人各主导一次后才重复。角色称号不限制技能。')
p(gm, '卡住时只问：“我们知道什么？还不知道什么？谁承担风险？下一步做什么？”每幕结束写一句决定即可。', True)

page(gm, '02 / 规则与公平约定', '开场用 60 秒说明。新增机制用于本次草案，不是课程评分规则。')
h(gm, '检定怎么结算')
p(gm, '玩家先说行动、目的和依据；GM 再公开 DC 与加成；玩家掷骰。默认一次只用最相关的 +3、+2 或 +1。普通的观察、读卡、说真话、提出伦理理由不要求掷骰。检定只处理有不确定性的准备或执行。')
p(gm, '以下把总点数减去 DC 的差额称为 Δ。正式表列出了 Δ≥3、0、−1/−2、−3/−4、≤−5；本草案明确补齐 Δ=+1/+2，按成功处理。')
outcomes(gm, [('Δ≥3｜成功并获益', '目标成功，另给该段列出的额外好处。'), ('Δ=0、+1、+2｜成功', '目标按计划完成。'), ('Δ=−1、−2｜接近成功', '目标大部分完成，伴随已说明的代价。'), ('Δ=−3、−4｜失败但有转机', '原行动未成，获得可继续行动的信息或备用条件。'), ('Δ≤−5｜失败并有后果', '产生较大但有限的损失；继续故事，不让玩家退出游戏。')])
h(gm, '应急物资：两次团队补救机会')
p(gm, '公会开场有 2 份应急物资。只有执行段能用：掷骰并看到结果后，小组可花 1 份，把该次结果提升一档；每次检定最多花 1 份，最高提升到“成功”，不能买到“成功并获益”。依序为：后果→转机→接近成功→成功。已达成功就不消耗。物资是额外绳索、电池、密封材料的抽象，不替代路线自带的基本设备。')
h(gm, '准备、创意与协作')
p(gm, '每项准备成功或成功并获益，对应执行 DC 降 2。玩家提出具体、合理的新方案，也可让一个相关执行 DC 降 2。两者取较高者，不叠加。帮忙者说明自己的贡献，GM 给叙事空间，不额外无限加点。不要求玩家熟悉真实工程计算；合理说明假设、咨询专业 NPC 即可。')
p(gm, '每个目标只检定一次。单纯“再试同一个动作”不重掷；改用备用手段体现在该档后果中。课堂讨论不扣游戏内秒数；三个执行目标对应三个潮浪窗口，不能通过拖延无限行动。')
h(gm, '难度边界')
p(gm, '准备检定 DC 11；执行基础 DC 为通路 12、转移 12、隔离 13。以 +3 为例，执行达标率为 60%、60%、55%；得到准备优势后为 70%、70%、65%。有创意不会被临时提高 DC 抵消。若教师允许叠加加成，本场会更容易；保留公开难度，不在掷骰后加难。')
p(gm, '坏骰不删除关键线索、不强制突然死亡，也不自动把欺骗变成正确决定。对真实风险的诚实说明与行动结果分别记录。总点数不代表课程分数或伦理优劣。')

page(gm, '03 / 第一幕：绿灯与海水', '0–4 分钟｜把人物放进危机，再给他们决定空间。')
say(gm, 'The Shell hums above the Luminis Sea. Inside, maps flicker, water turbines spin, and somebody has left seven Swiss Army knives on the same table. Outside, your Maid-Mer guard drops an evil seagull beside the door. Then the alarm begins. The coastal pump station is flooding. Twelve maintenance workers are inside. On your screen, one message keeps blinking: GREEN — SAFE TO OPERATE.')
say(gm, 'The next storm surge will reach the station in twelve minutes. You have three action windows once your plan is ready. Your job is to open a rescue route, move the workers, and isolate the station safely. You also need to tell the other guilds what is actually happening. What do you do first?')
h(gm, '先给所有人知道的事实')
bullet(gm, '被困的是外围泵站的 12 名维修人员，不是四艘 Nemorath 运输船上的乘客。本故事不改变老师已经决定的运输船结果。')
bullet(gm, '泵站与 The Shell 相邻但独立；关停会让总部实验室停用到次日，诊所有独立备用电源。十二分钟是剧情节奏，不是要求玩家真实倒计时。')
bullet(gm, '泵站上层有干燥避难室。若撤离不完整，仍可临时容纳 4 人约两小时；备用救援班要等暴潮减弱后才能到。先公开这个保底条件。')
bullet(gm, '基本设备都可取用：安全绳、手摇绞盘、额定载 8 人的小工作艇、两名 NPC 船员、可手动操作的外部闸门。两份应急物资另算。')
h(gm, 'GM 私下知道的真相')
p(gm, '绿灯来自 17:10 的低负荷测试；现在是 17:40，预计 17:52 暴潮抵达。测试没有覆盖暴潮负荷，因此不能证明当前安全。海鸥损坏了状态回传线，屏幕冻结显示旧值；本地机械指示仍可读。Mira 批准过临时旁路，却没有完成满负荷验证。她害怕失去公会信任，倾向先保住设备再解释；她不是蓄意害人的反派。')
p(gm, '玩家无需找出“谁坏”。目标是验证证据、组织救援、识别自己的能力边界，并决定如何处理不完整记录和承诺。关键线索全部在玩家证据页上。')
h(gm, '幽默的用法')
p(gm, '七把瑞士军刀、恶海鸥和 Maid-Mer 是已有组内故事的回扣；可用一句笑话缓解紧张。Maid-Mer 为鱼身人腿，主要通过跺脚、摆头和丢海鸥表达。它不突然讲人话，也不会替大家修好泵站。若老师已改变其关系，改成远处的熟悉叫声，不强行召回盟友。')

page(gm, '04 / 第二幕：三张卡，一幅真相', '4–13 分钟｜推理 + 三次不同玩家主导的准备行动。')
p(gm, '把证据页分给三个玩家或三组。每人先用 20–30 秒转述自己的卡，再用 90 秒共同回答：“绿灯证明了什么？没有证明什么？哪件事应该立刻做？”不设唯一措辞；抓住旧数据、机械现场证据和救援条件即可。')
h(gm, '证据 A / 冻住的绿灯')
p(gm, '显示：SAFE，最后回传 17:10，低负荷测试。纸面变更记录：临时旁路，满负荷验证待做。现在 17:40。结论：存在未经核实的风险；不能直接推断一定安全，也不能直接断言任何人故意造假。')
h(gm, '证据 B / 三条可走的路')
p(gm, '崖壁上有经目测尚完整的锚点；海面背风侧有短暂平水区，工作艇可两趟各运 6 人；外部泄洪渠可降低泵站周围水位，但会影响一片本地育幼浅滩。每条路线都有设备可用；任何一张卡都不锁住某一条路线。')
h(gm, '证据 C / 十二个人与一份记录')
p(gm, '12 人中有 2 人行动不便，均能在协助下移动；上层避难室可临时容纳 4 人。诊所备用电源可维持当前关键设备 20 分钟，救援结束前需完成交接；停泵不会立即造成医疗中断。报告模板的“适用条件、未验证部分、责任人”栏仍空白。')
h(gm, '三次准备检定：保证三位不同玩家主导')
outcomes(gm, [('① 结构与路线 / DC 11', '如何验证锚点、潮流或外部通路？相关工程技能、具体工具或团队组织均可。达标：勾“通路准备”，执行通路 DC 12→10。'), ('② 人员与交接 / DC 11', '如何确认人数、行动需求、两趟运载或医疗交接？达标：勾“转移准备”，执行转移 DC 12→10。'), ('③ 能源与证据 / DC 11', '如何区分旧回传、本地状态和关键供电？达标：勾“隔离准备”，执行隔离 DC 13→11。')])
p(gm, '三个准备检定使用同一结果逻辑：成功并获益＝优势 + 额外一条证据/一位 NPC 主动配合；成功＝优势；接近成功＝线索全给，但尚未验证到足以降低 DC；失败但有转机＝尝试受阻，NPC 提供备用方法，线索仍全给；失败并有后果＝验证工具暂时失灵，仍给线索，后续依靠基本设备，执行 DC 保持原值，不再追加罚点。')
say(gm, 'The green light is a claim. The water on the floor is evidence. Tell me what you would verify, who can help, and what you will do if your first method fails.')
p(gm, '若玩家直接提出有效推理，肯定其推理；掷骰处理的是实地准备质量，不用于决定推理是否“被允许正确”。', True)

page(gm, '05 / 第三幕：有压力的人，不同的利益', '13–20 分钟｜角色扮演与协商；伦理理由不靠骰子判定。')
h(gm, 'Mira / Builders 的临时工程负责人')
say(gm, 'I approved the temporary bypass. The low-load test passed, but I did not finish the surge test. If we abandon the station, we lose the spare motor. Could you sign the green status so the other guilds do not panic?')
p(gm, '她想保住人员、设备和声誉。面对尊重而具体的提问，她会承认记录不完整，并提供结构知识。玩家可以拒绝无条件签字，同时提出带适用范围的真实通报。她不会因为被拒绝就锁门或拒绝救人。')
h(gm, 'Tavi / Linkers 的通信协调员')
say(gm, 'I have an open channel to the other guilds. Give me three things: what you know, what you do not know, and what you need people to do. I can transmit a warning now and update it later.')
p(gm, 'Tavi 接受不确定性，希望避免“安全/危险”二选一的假精确。无需掷骰就能发事实通告。若玩家想用煽动、威胁或隐瞒换服从，明确记录其说法与可能破坏的信任；不增加奖励数值。')
h(gm, 'Dr. Vale / Life Weavers 的临床联络人')
say(gm, 'Two workers need assisted movement. My backup power lasts twenty minutes. Give me the handover time and a named contact. Do not promise continuous power unless you have verified it.')
p(gm, '她重视需求和可执行的交接，不按身份给生命排序。玩家一旦安排联系人及备用电源，她就配合；不另收一个“说服医生”的检定。')
h(gm, '给玩家的协商问题')
for t in ['我们确定了什么，还需要哪位具备相关能力的人核实？', '哪个方案把风险转嫁给了不在现场的人或生命？', '我们能承诺什么，不能承诺什么？', '谁负责记录和复查？如果负责人意见被否决且仍有安全风险，应该通知谁？']:
    bullet(gm, t)
p(gm, '真实通告示例：', True)
say(gm, 'The pump station is not verified for the current surge. Twelve workers are being evacuated. The Shell labs will go offline; clinic backup is active. Keep the coast clear. Mira will verify the isolation, and we will update you after the third action window.')
p(gm, '玩家愿意承担责任、如实说明未知、邀请专业人员核验，都是可见的领导力表现。不要根据是否背出 NSPE 条文给剧情加分。')

page(gm, '06 / 第四幕：选路线，先看代价', '20–24 分钟｜每条路线同样三次检定：通路 12、转移 12、隔离 13。')
h(gm, 'A / 崖壁救援：The High Line')
p(gm, '操作：固定高处锚点，用安全绳和手摇绞盘把人员接到 The Shell 上方平台，再从外侧隔离泵站。基本设备可用，准备优势也可用。')
p(gm, '已知代价：狭窄崖道只够运人，暴潮前带不走沉重的备用电机。电机留在将被淹的下层，后续需要检修或更换。生态浅滩不主动受扰；总部实验室仍须停用到次日。')
p(gm, '适合：机械、结构、组织、照护；不要求会攀岩的角色。提醒：两名行动不便者也有吊带，不因角色身份增加 DC。')
h(gm, 'B / 海上转运：The Quiet Water')
p(gm, '操作：以背风平水区为通道，用额定载 8 人的工作艇撤离；两名 NPC 船员占两席，每趟运 6 人，共两趟。海侧检修口可把备用电机锁上高位架避水，不占船上席位；此动作纳入通路安排。一次“转移”检定结算两趟协作，不逐人掷骰。')
p(gm, '已知代价：耗尽并损坏一组实验室移动电池，需要更换；救援艇甲板空间全部给人，低层工具箱留下。诊所独立备用电源不被拿走，需按时交接。电机可保住，代价转移到便携能源和工具。')
p(gm, '适合：导航、通讯、人员调度、环境观察。若玩家想让 Maid-Mer 指出平水区，接受创意降 DC，但与已有准备优势不叠加。')
h(gm, 'C / 外部泄洪：The Borrowed Channel')
p(gm, '操作：手动分流降低外侧水压，沿检修步道接人，再隔离泵站。不得把人派到仍带电的进水间；阀门位于外侧可达位置。')
p(gm, '已知代价：水会通过一片育幼浅滩，扰动一季的部分栖息地；总部实验室停用到次日，电机被泥沙污染，需要清理。玩家需提出尽量减轻生态影响和事后恢复的方案。')
p(gm, '适合：水利、环境、手工操作、跨组协调。保护人命的紧迫性与生态修复责任都可进入讨论；不把浅滩损伤当作隐瞒的惩罚。')
h(gm, 'D / 混合或自创方案')
p(gm, '允许。让玩家说明通路、人员、隔离三个目标分别怎样达成；仍是三次检定，基础 DC 不增加。若具体创意减少一种已知代价，允许它奏效，并把设备容量、剩余风险或修复义务讲清楚。不要为了强迫二选一突然制造新受害者。每个相关目标最多降 DC 2。')
p(gm, '小组用 60 秒决定：选什么；接受什么代价；保障谁；什么条件会让我们改变计划。允许一人保留异议。意见不一致时，先复述共同目标，再比较证据，不额外扣时间。', True)

page(gm, '07 / 第五幕：第一浪，建立通路', '24–29 分钟｜执行①｜基础 DC 12；准备/创意有效时 DC 10。')
say(gm, 'The first wave hits the pump station. The green light stays green. A seagull lands on the alarm speaker and screams louder than the alarm. Your route must be ready before the next wave. Who is leading, and how is everyone helping?')
p(gm, '从还没主导检定的人中邀请一位。先让其说行动和依据，再宣布最终 DC。其余人各说一句贡献。崖壁是锚点与绞盘，海上是系泊与两趟航线，泄洪是分流与步道，均使用同样的结果梯度。')
outcomes(gm, [('成功并获益', '通路稳固开放。顺手取回一件个人工具或一份故障记录，玩家决定是哪件；不再增加检定加成。'), ('成功', '通路按计划开放，第二浪可以开始转移。'), ('接近成功', '通路开放；一件可替换的公共装备损坏。说清是哪件，不让任何玩家永久失去唯一工具加成。'), ('失败但有转机', '原通路未能完全建立；NPC 展开预置的慢速手动备用通路，仍进入转移段。代价是附近普通物资被潮水卷走；玩家保住了标记安全线的方法。'), ('失败并有后果', '主通路受损；改用狭窄备用通路。下层工棚报废，需要其他公会未来协助修复。转移段仍可进行，不追加 DC 或伤害罚点。')])
h(gm, '坏骰之后给一个明确选择')
say(gm, 'You can accept that consequence, or spend one emergency kit to improve this result by one step. You have no more than two kits for the entire rescue. What does the group choose?')
p(gm, '提升后使用新档完整叙述，原档后果被替换；不要把两档惩罚同时留下。A/B/C 路线本来公开的代价仍保留。')
h(gm, '记录与转场')
p(gm, '记录通路结果：________。剩余物资：□ □。只记实际留下的损失；本段已使用的准备优势划掉。')
say(gm, 'The route is usable. It may not be the route you wanted, but your people can now move. The twelve workers are waiting. Two need assistance, and the next wave is already turning the water white.')
p(gm, '若玩家先隔离再通路，允许重排三目标；确认其绞盘改用手摇或独立电池即可。每个目标仍只结算一次。演出顺序可变，公开目标和难度不变。')

page(gm, '08 / 第二浪与第三浪', '29–40 分钟｜执行②转移 DC 12；执行③隔离 DC 13。准备/创意分别最多降 2。')
h(gm, '第二浪：把人接回来')
say(gm, 'You hear twelve voices through the radio. Some are calm. Some are talking over each other. Your plan needs a movement order, assistance for two workers, and a headcount at the receiving point. Who leads this part?')
outcomes(gm, [('成功并获益', '12 人全部撤离，连同一箱基本医疗物资一起带回。'), ('成功', '12 人全部安全撤离并完成清点。'), ('接近成功', '12 人全部撤离，2 人需观察；明确后续医疗交接，人员不退出游戏。'), ('失败但有转机', '8 人撤离，4 人进入上层干燥避难室；双向无线电保持，外部备用救援班已获位置。未撤离者的获救需要暴潮后的跟进任务。'), ('失败并有后果', '8 人撤离，4 人进入同一避难室；无线电失效，但 Tavi 保留人员名单和位置。已有两小时临时保障，后续优先恢复联系。')])
p(gm, '组员可按规则花 1 份物资提升一档。不要在后续隔离失败时把已经撤离的人重新“杀掉”。记录：已撤离____；避难室____；医疗/联系安排____。')
h(gm, '第三浪：让危险停止，并留下真记录')
say(gm, 'The local indicator is still live. The remote green light cannot confirm isolation. Mira is ready to help verify a manual shutdown. Tavi is holding the channel open. Tell us who checks the system, who checks the backup, and what you will report.')
p(gm, '玩家明确检查本地状态、诊所交接和记录责任后再掷执行③；这些事项可以由不同玩家承担。隔离失败不取消避难室的独立临时保障。')
outcomes(gm, [('成功并获益', '本地隔离被验证；诊所交接完成；保存原始日志，并找到可复用的备件。'), ('成功', '隔离和备用供电交接完成；按玩家实际选择留下真实或不完整的报告。'), ('接近成功', '隔离完成，但留下需次日复核的临时封堵；具名负责人接手，实验室继续停用。'), ('失败但有转机', '隔离状态尚未验证；外围封锁防止进入，独立备用供电正常。机械指示提供了下次检修位置。'), ('失败并有后果', '泵站仍需封锁，部分设备被淹；诊所继续使用已知备用电源，公会必须在 20 分钟保障结束前协调后续供电。追加停电、维修任务；不凭空宣布病人死亡。')])
p(gm, '最后一次补救仍可用。无论骰点多高，“没验证却写成已验证”都仍是玩家作出的报告选择。反过来，诚实报告不会因低骰被改写成撒谎。')

page(gm, '09 / 结局、变化与简版', '40–45 分钟｜逐项结算，不把一切压成一个道德分数。')
h(gm, '结局按三条线分别说')
p(gm, '先说人员：12 人安全 / 12 人撤离但需观察 / 8 人撤离且 4 人待后续救援。再说设施：已隔离 / 临时隔离需复核 / 封锁待处理。最后说信任：如实披露了什么、隐瞒了什么、谁承诺跟进。再补上路线公开代价与实际损失。一个成功不自动抹去另一个失败。')
say(gm, 'The last wave breaks against the cliffs. The Shell still stands. Now tell the guilds: how many people are accounted for, what remains unsafe, and who owns the next action. The story ends with that responsibility, not with the green light.')
h(gm, '三种可用尾声：按事实改写')
bullet(gm, '“灯灭了，承诺留下了”：人员全撤、隔离完成、报告真实。Mira 在旧日志旁签上改正说明；新挂牌写着适用条件、验证人和复查时间。路线代价进入修复清单。')
bullet(gm, '“守住这一浪”：人员全撤但设施仍封锁。Tavi 广播临时限制；公会准备下一轮维修和备用供电。胜利包含尚未完成的工作。')
bullet(gm, '“四个名字还在频道里”：8 人已撤，4 人在避难室。把两小时临时保障、位置、联系状况和后续负责人明确给玩家；下次任务从这份交接开始。')
h(gm, '只在还有时间时加一个变化')
p(gm, '可选 ① 恶海鸥叼走挂牌：玩家用任意两件现有物品提出替代标记，30 秒，不掷骰。② Maid-Mer 拒绝帮忙，直到有人说明将如何恢复浅滩；真诚讨论即可。③ 一名工人自告奋勇最后走；小组讨论自愿与风险告知。不再添加检定、资源扣除或突发死亡。')
h(gm, '20 分钟简版 + 5 分钟复盘')
p(gm, '0–3 分开场并发三卡；3–7 分各自转述并选择路线；7–9 分做三名 NPC 的简短问答；9–17 分做三次执行检定；17–20 分结局与口头报告。跳过准备检定和准备优势，保留两份应急物资与创意最多降 DC 2。基础 DC 仍为 12/12/13。')
p(gm, '六位玩家两人一组：一人提出行动，另一人补充依据和备用方案，再由其中一人掷骰。三组分别主导三个目标，确保六人都发言。额外 5 分钟用下一页三道核心复盘题。')
h(gm, '人数与节奏调整')
p(gm, '4–5 位玩家：合并调查小组，轮换主导；不增加目标。7–8 位玩家：每次检定增加一名明确的顾问/记录者，不增加叠加加成；决策前先每人一句。还剩 10 分钟却未执行：立刻结束 NPC 问答，依次完成三个目标，再用一句话交接。')

page(gm, '10 / 复盘与 Co-GM 操作清单', '结局之后退出角色：先个人反思，再谈小组。')
h(gm, '十分钟复盘')
bullet(gm, '1 分钟静写：我刚才做的一个决定、当时依据、现在最不确定的一点。')
bullet(gm, '3 分钟轮流：绿灯、现场海水、测试日志中，我们最先相信哪个？为什么？')
bullet(gm, '3 分钟讨论：谁承担了我们方案的代价？我们怎样处理公众安全、能力边界和真实通报？')
bullet(gm, '2 分钟迁移：下一次真实工程团队遇到未验证状态，我们会建立什么验证与升级通报流程？')
bullet(gm, '1 分钟收尾：每人说一件队友做得好的具体事情；两位 Co-GM 各记录一个想改进的主持动作。')
p(gm, '5 分钟简版只问三题：① 我们依赖了什么证据？② 谁承担代价、谁没被听见？③ 下一次工程协作要改变哪一个流程？')
h(gm, '把课程目标变成可见的动作')
bullet(gm, '创新：给自创方案同等处理机会；表扬具体因果，不要求玩家猜 GM 想法。')
bullet(gm, '全员参与：用六人主导名单；每次请还没说话的人先补充，允许跳过表演、改用读卡或记录。')
bullet(gm, '清楚指导：每幕重述目标、当前已知事实、允许的下一步；掷骰前公开 DC。')
bullet(gm, '放心表达：邀请不同意见；针对行动提问，不拿现实身份、口音或能力开玩笑。')
bullet(gm, 'Co-GM 协作：准备时各试读一段；行动段交换职责；规则分歧先用公开约定完成本次，课后再核查。')
h(gm, '常见岔路怎么接')
outcomes(gm, [('玩家想战斗', '可阻挡海鸥、拉开碎片、护住绳线；将其变成通路/转移目标中的动作。暴力不能破坏未知故障后自动变安全，不额外展开 HP 战。'), ('玩家想叫全体公会解决', '通信能获得专业建议和事后援助；近场行动由 Guardians 完成。备用班抵达时间在开场已公开，不临时拖延。'), ('玩家要等到完全确定', '说明可以维持封锁、使用避难室并交接；这是可玩的选择。不要把所有等待都判定为道德失败。'), ('玩家发现更好的方案', '接受合理创新。对承重、人数、能量等关键假设请专业 NPC 提供简明边界；不要求真实专业知识，不事后变更事实。'), ('所有骰子都差', '依然得到线索、8 人撤离及4人临时保障、封锁和后续任务。选择是否用物资改善结果，保留可继续玩的故事。')])
p(gm, '桌面记录：路线____；准备□通路 □转移 □隔离；物资起始2 / 剩余____；通路后果____；人员状态____；隔离状态____；报告原话____；负责人/复查时间____。')

page(gm, '11 / 材料依据与改编边界', '供两位 GM 备课时查看；不必在游戏中朗读。')
h(gm, '直接沿用的背景')
bullet(gm, 'GUardians of the Galaxy HQ.docx：The Shell、Luminis Sea、崖壁与热带海洋、风水能源、海鸥和鱼身人腿的生物。')
bullet(gm, 'Guardian Guild Mission Statement.docx：保护其他生命、忠诚、积极沟通、开放协调。')
bullet(gm, 'GM WEEK 1.docx：Maid-Mer 遭遇、玩家已经选择友善接触、七把瑞士军刀的组内趣事；“Yes, and”、记录讨论与推进活动。')
bullet(gm, 'Nemorath Communication & Guild Prioritization Task.docx：避难者安置、水与基础设施、跨公会合作、避免冲突。')
bullet(gm, '第6周课件：基础设施仓促建成、医疗资源紧张、Life Weavers 与 Builders 的临时医院。本文没有推翻老师对运输船或其他公会的既定结果。')
h(gm, '课程原规则与学习主题')
p(gm, 'Game Rules & Mechanics.pdf，第1–3页：D20、相关加成、GM设置DC、按差额描述五档后果、鼓励创意、避免“什么也没发生”。课程说明游戏时通常不使用手机和电脑，所以本草案可完全靠打印材料运行。')
p(gm, 'LMS 第8周 Choosing a[n] (Ethical?) Path：工程伦理、困难决策与过去行为的影响。NSPE课程PDF：I.1 公众安全，I.2 能力范围，I.3 真实客观通报，III.1.a 承认错误，III.2.d 可持续发展。这里是课堂情境应用，不声称完成任何工程标准认证。')
p(gm, 'Syllabus，第2页：Co-GM引导与协作；第4页：自评20 + 同伴评价20。实际评价表关注创新、参与、指导、表达舒适度、Co-GM合作及领导力特征。本文游戏结果不换算成课程分数。')
h(gm, '为这份草案新写的内容')
p(gm, '泵站、十二名维修人员、三名NPC、旧绿灯的故障原因、时间、避难室和医疗备用条件、三条路线及其代价、固定DC、单一加成约定、两份物资和准备优势均为原创场景设计。没有把 Jacob and Gautam 的五事件 campaign 当作老师要求，也没有复制其 Boss/Readiness 系统。')
h(gm, '使用前核对')
p(gm, '确认教师实际给本组的带团时长、当前剧情进展和加成叠加方式。Syllabus 第5页规定课程作业使用 AI 需教师事先批准；本文件是明确标注的 AI 辅助准备草案，作为课堂成果使用时需遵循该要求。')
p(gm, '课程入口：https://lms.rpi.edu/ultra/courses/_26943_1/outline')
p(gm, '第8周：https://lms.rpi.edu/ultra/courses/_26943_1/document/_1028137_1?view=content&state=view')
p(gm, 'Co-GM评价：https://surveys.rpi.edu/jfe/form/SV_4UXvqLDBSGMnl9s')

pack = make_doc('Under the Green Light | Player Pack', 'Printable English briefing, evidence, stakeholder prompts and report')
pack.add_heading('UNDER THE GREEN LIGHT', 0)
p(pack, 'PLAYER BRIEFING  /  GUARDIANS GUILD', True)
p(pack, 'The Shell, Luminis Sea, Heliora Prime')
h(pack, 'The situation')
p(pack, 'The coastal pump station is flooding. Twelve maintenance workers are inside. Your screen says GREEN — SAFE TO OPERATE. A storm surge will arrive in twelve fictional minutes. Once your plan is ready, you have three action windows. Discussion time does not secretly consume those windows.')
h(pack, 'Your three objectives')
for t in ['OPEN A ROUTE. Build or secure a usable rescue path.', 'MOVE THE PEOPLE. Account for twelve workers; two need assisted movement.', 'ISOLATE THE HAZARD. Verify the local shutdown and medical power handover.']:
    bullet(pack, t)
h(pack, 'Known safeguards and equipment')
p(pack, 'An upper refuge can hold four people for roughly two hours if evacuation is incomplete. The outside rescue crew arrives after the surge eases. The clinic has independent backup for twenty minutes. The station and refuge safeguards are separate. Shutdown affects noncritical Shell labs until the next day; it does not instantly cut clinic power.')
p(pack, 'Available: safety lines, harnesses, a hand winch, an eight-person-rated workboat, two NPC boat crew, and an exterior manual gate. The boat has six evacuee seats after its two crew; two trips move twelve people. You also have TWO shared emergency kits.')
h(pack, 'Choices you can propose')
bullet(pack, 'HIGH LINE: cliff rescue. Known cost: leave the heavy spare motor below the flood line for later repair or replacement.')
bullet(pack, 'QUIET WATER: two boat trips; secure the motor on its high rack from the sea-side hatch. Known cost: replace the spent lab battery pack and leave low-level tool crates behind. Clinic backup stays separate.')
bullet(pack, 'BORROWED CHANNEL: divert the water, then use the walkway. Known cost: disturb part of a nursery shallows for a season; commit to restoration and clean the motor.')
bullet(pack, 'YOUR PLAN: combine or replace these approaches. Explain access, people, isolation, and the assumptions your plan depends on.')
h(pack, 'This draft’s session conventions')
p(pack, 'D20 + one justified bonus: Leadership +3 OR Engineering +2 OR Tool +1. Preparation or a concrete creative solution may lower the relevant execution DC by up to 2; these reductions do not stack. Base execution DCs: route 12, people 12, isolation 13. Teacher instructions take priority.')
p(pack, 'After seeing an execution result, the group may spend one emergency kit to improve it by one result tier, at most once per roll and never above ordinary success. There are no extra natural-1 penalties. Say what you try; the GMs explain the result.')

page(pack, 'EVIDENCE CARDS', 'Cut along the spaces if useful. Share each card aloud; nobody has to solve the mystery alone.')
h(pack, 'CARD A / THE GREEN LIGHT')
p(pack, 'CONTROL SCREEN: SAFE TO OPERATE\nLast received status: 17:10\nTest condition: low load\nCurrent time: 17:40\nExpected surge: 17:52\nChange log: temporary bypass installed; full-surge verification PENDING.')
p(pack, 'A gull-damaged return cable cannot send a fresh status. The local mechanical indicator is still readable. Water is entering the station now.')
p(pack, 'Discuss: What does the old test establish? What does it fail to establish? What must you verify locally?', True)
p(pack, '________________________________________________________________')
h(pack, 'CARD B / THE THREE ROUTES')
p(pack, 'CLIFF: an anchor point and a hand winch are available; check their setup. SEA: a short lee-side calm window supports two workboat trips. CHANNEL: an exterior manual gate can reduce water around the walkway; runoff crosses nursery shallows.')
p(pack, 'All three routes have basic equipment available. None requires a specific character class. A useful plan can combine routes. There is no hidden fourth route you must guess.')
p(pack, 'Discuss: What can go wrong? What would make the route usable? Which consequence are you willing to accept?', True)
p(pack, '________________________________________________________________')
h(pack, 'CARD C / PEOPLE AND HANDOVER')
p(pack, 'Twelve workers; two need assisted movement. Four can wait in the upper dry refuge for roughly two hours. Clinic backup has twenty minutes. A named person must coordinate power handover before it ends. The backup rescue crew has been told when the surge window closes.')
p(pack, 'The incident-report fields for operating limits, unverified conditions, verification owner, and next update are blank.')
p(pack, 'Discuss: How will you count people? Who handles the medical handover? What can you responsibly promise?', True)
p(pack, '________________________________________________________________')

page(pack, 'STAKEHOLDER PROMPTS', 'Optional: three players can voice these concerns during discussion, then return to their own characters. No secret objectives or player-versus-player rewards.')
h(pack, 'MIRA / BUILDERS')
p(pack, 'You approved the temporary bypass. The low-load test passed; the full-surge test was not completed. You want the workers safe, but you also fear losing the spare motor and everyone’s confidence.')
p(pack, 'Opening: “Could we keep the green status while we work? I am afraid the other guilds will stop trusting our equipment.”')
p(pack, 'You can offer structural knowledge and practical help. If the group asks for honest limits and a correction, cooperate. You do not withhold rescue help when they refuse an unrestricted safety statement.')
p(pack, 'Question for the group: “What exactly can we certify, and who verifies it?”', True)
p(pack, '________________________________________________________________')
h(pack, 'TAVI / LINKERS')
p(pack, 'You have an open communication channel. Other guilds need usable instructions, not certainty that the evidence cannot support.')
p(pack, 'Opening: “Give me what we know, what we do not know, and what people should do now. I can update the message later.”')
p(pack, 'You can transmit the location, hazards, headcount, restrictions and next update. You preserve the original message and record later corrections.')
p(pack, 'Question for the group: “Who needs this information, and when will we update them?”', True)
p(pack, '________________________________________________________________')
h(pack, 'DR. VALE / LIFE WEAVERS')
p(pack, 'Two workers need assistance. You have independent clinic backup for twenty minutes. You need a realistic arrival plan and a named power-handover contact.')
p(pack, 'Opening: “Tell me who needs assistance, who is bringing them, and who will confirm my next power source.”')
p(pack, 'You help once the group gives workable arrangements. You do not rank lives by status or make rescue conditional on a persuasion roll.')
p(pack, 'Question for the group: “Whose needs are missing from our plan?”', True)

page(pack, 'EVERYONE HAS A JOB', 'Six temporary focus roles. Keep your own character; choose or swap a focus. No focus grants extra numerical bonuses or exclusive access.')
for title, text in [
    ('1 / STRUCTURES & ACCESS', 'Ask what supports the route. Explain a check of anchors, walkway, mooring or barriers. Lead one preparation or execution action.'),
    ('2 / WATER & ENVIRONMENT', 'Observe the tide, calm water and nursery shallows. Identify a route cost and a way to reduce harm.'),
    ('3 / PEOPLE & CARE', 'Plan assisted movement and a headcount. Ask who receives the workers and checks the medical handover.'),
    ('4 / POWER & TOOLS', 'Distinguish a screen claim from verified local state. Consult the appropriate expert before confirming isolation.'),
    ('5 / COMMUNICATION & TRUST', 'Draft a short factual message. Invite one concern the group has not heard. Name an owner and a next update.'),
    ('6 / COORDINATION & CHALLENGE', 'Summarize the plan, capacity limits and fallback. Ask one respectful “what if?” before the group commits.')
]:
    h(pack, title)
    p(pack, text)
h(pack, 'Before each roll')
p(pack, '“I will __________ so that __________. My relevant skill, competency or tool is __________. __________ will help by __________. If this fails, our fallback is __________.”')
h(pack, 'Make a group choice')
p(pack, 'Our route: ____________________________________________________\nOur three action leads: ___________________________________________\nA consequence we knowingly accept: ________________________________\nOne assumption to verify: _________________________________________\nOne voice or need we have not yet included: __________________________')
p(pack, 'A different opinion is useful information. You can speak briefly, point to a card, or ask someone to read your idea. Performance and accents are optional.')

page(pack, 'INCIDENT REPORT & DEBRIEF', 'Complete in brief phrases or give a 30-second spoken handover. This is an in-game worksheet, not a claim of engineering certification.')
for title, prompt in [
    ('1 / WHAT WE KNOW', 'Verified observations and current headcount:'),
    ('2 / WHAT REMAINS UNCERTAIN', 'Unverified equipment, people waiting, medical or environmental follow-up:'),
    ('3 / OUR DECISION', 'Route, reason, accepted consequences, and any change of plan:'),
    ('4 / NEXT ACTION', 'Named owner, restrictions, backup arrangements, and next update time:'),
]:
    h(pack, title)
    p(pack, prompt)
    p(pack, '________________________________________________________________\n________________________________________________________________')
h(pack, 'After the story: step out of character')
bullet(pack, 'Which evidence did we trust first? What changed our minds?')
bullet(pack, 'Who bore the cost of our plan? Whose voice did we miss?')
bullet(pack, 'How did we handle safety, competence, truthfulness, and correction of errors?')
bullet(pack, 'What one process would we use in a real engineering team next time?')
p(pack, 'One specific contribution I appreciated: ______________________________\nOne thing I would change about my own teamwork: ______________________')
p(pack, 'This fictional module adds the pump station, local incident, NPCs and session conventions to the class world. Apply your instructor’s current rules and timeline. AI-assisted preparation draft; course use follows the syllabus requirements.')

gm.save(OUT / 'Under the Green Light - GM Script.docx')
pack.save(OUT / 'Under the Green Light - Player Pack.docx')
print('Created', OUT / 'Under the Green Light - GM Script.docx')
print('Created', OUT / 'Under the Green Light - Player Pack.docx')
